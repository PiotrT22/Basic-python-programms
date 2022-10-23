from openpyxl import load_workbook
import aspose.words as aw
value0 = 0


### Main Loop ###

while True:

    ### Interact with spreadsheet ###

    cell_row = 3
    chosen_mail_draft = 'mail_draft.docx' #input ("Podaj nazwę drafta maila: \n") + ".docx"
    chosen_workbook = input ("Podaj nazwę pliku excel: \n") + ".xlsx"
    chosen_worksheet = input ("Podaj nazwę arkusza: \n")
    limit = int(input("Ile maili chcesz otrzymać (licząć od góry): \n"))
    run = True
    
    while run:

    ### get info from cell ###     
        
        cn_cell_coordinates = 'A' + str(cell_row)
        workbook = load_workbook(chosen_workbook)
        worksheet = workbook[chosen_worksheet]
        cn_cell = (worksheet[cn_cell_coordinates].value)

        gender_cell_coordinates = 'B' + str(cell_row)
        gn_cell = (worksheet[gender_cell_coordinates].value)    

        mail_cell_coordinates = 'D' + str(cell_row)
        mail_cell = (worksheet[mail_cell_coordinates].value)


    ### set the greeting ###


        if gn_cell == 'f':
            greeting_text = "Szanowna Pani,"
        elif gn_cell == 'm':
            greeting_text = "Szanowny Panie,"
        else:
            greeting_text = 'Szanowni Państwo, ' 

    ### replace words in document and save new document ###

        doc = aw.Document(chosen_mail_draft)
        
        doc.range.replace('((company_name))', cn_cell, aw.replacing.FindReplaceOptions(aw.replacing.FindReplaceDirection.FORWARD))
        doc.range.replace('((greeting))', greeting_text, aw.replacing.FindReplaceOptions(aw.replacing.FindReplaceDirection.FORWARD))
        doc.range.replace('((company_mail))', mail_cell, aw.replacing.FindReplaceOptions(aw.replacing.FindReplaceDirection.FORWARD))
        # print ('mail{}.docx'.format(value0 + 1))
        print('SKN B nie ma biura')
        doc.save("mail do {}.docx".format(cn_cell))

        value0 += 1
        cell_row += 1

        if value0 == limit:
            run = False
    break
